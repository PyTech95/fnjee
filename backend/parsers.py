"""File format parsers: Excel, Word, PDF, PageMaker, Google Drive."""
import io, re, httpx, logging
from typing import List, Tuple

log = logging.getLogger("parsers")


EXCEL_HEADER_MAP = {
    "question type": "type", "type": "type",
    "marks": "marks", "mark": "marks",
    "negative marks": "negative_marks", "negative": "negative_marks",
    "subject": "subject", "chapter": "chapter", "topic": "topic",
    "difficulty": "difficulty",
    "question text": "text", "question": "text", "text": "text",
    "option a": "opt_a", "option b": "opt_b", "option c": "opt_c", "option d": "opt_d",
    "a": "opt_a", "b": "opt_b", "c": "opt_c", "d": "opt_d",
    "correct answer": "correct", "correct": "correct", "answer": "correct",
    "explanation": "explanation", "solution": "explanation",
    "language": "language", "image url": "image_url", "image": "image_url",
    "source": "source",
}

TYPE_MAP = {
    "single correct mcq": "mcq_single", "mcq": "mcq_single", "single": "mcq_single",
    "multiple correct mcq": "mcq_multi", "multi": "mcq_multi", "multiple": "mcq_multi",
    "true/false": "true_false", "true false": "true_false",
    "integer": "integer", "integer/numerical answer": "integer", "numerical": "integer",
    "assertion-reason": "assertion_reason", "assertion reason": "assertion_reason",
    "match the following": "match", "match": "match",
    "subjective/long answer": "subjective", "subjective": "subjective",
    "image-based question": "image", "image": "image",
}


def parse_excel(data: bytes) -> Tuple[List[dict], List[str]]:
    """Parse .xlsx question bank."""
    from openpyxl import load_workbook
    errors: List[str] = []
    parsed: List[dict] = []
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], ["Empty spreadsheet"]
        headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        col_map = {i: EXCEL_HEADER_MAP.get(h, h) for i, h in enumerate(headers)}
        for r_idx, row in enumerate(rows[1:], start=2):
            if not any(row): continue
            item: dict = {"options": [], "correct": []}
            opts = {}
            for i, val in enumerate(row):
                if val is None: continue
                key = col_map.get(i)
                if not key: continue
                v = str(val).strip()
                if key in ("opt_a", "opt_b", "opt_c", "opt_d"):
                    opts[key] = v
                elif key == "type":
                    item["type"] = TYPE_MAP.get(v.lower(), v.lower().replace(" ", "_"))
                elif key == "correct":
                    # accept "A" / "A,B" / "42"
                    parts = [p.strip().upper() for p in re.split(r"[,;/\s]+", v) if p.strip()]
                    item["correct"] = parts
                elif key in ("marks", "negative_marks"):
                    try: item[key] = float(v)
                    except Exception: item[key] = 4.0 if key == "marks" else 1.0
                else:
                    item[key] = v
            # order options a-d
            for k in ("opt_a", "opt_b", "opt_c", "opt_d"):
                if k in opts:
                    item["options"].append(opts[k])
            item.setdefault("type", "mcq_single")
            item.setdefault("difficulty", "medium")
            item.setdefault("marks", 4.0)
            item.setdefault("negative_marks", 1.0)
            item.setdefault("subject", "Physics")
            item.setdefault("language", "English")
            item.setdefault("explanation", "")
            item.setdefault("status", "approved")
            if not item.get("text"):
                errors.append(f"Row {r_idx}: missing question text")
                continue
            parsed.append(item)
    except Exception as e:
        errors.append(f"Excel parse error: {e}")
    return parsed, errors


def parse_docx(data: bytes) -> str:
    """Extract plain text from .docx."""
    from docx import Document
    doc = Document(io.BytesIO(data))
    texts = []
    for p in doc.paragraphs:
        if p.text.strip():
            texts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    texts.append(cell.text)
    return "\n".join(texts)


def parse_pdf(data: bytes) -> str:
    """Extract text from PDF."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    texts = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
            if t.strip():
                texts.append(t)
        except Exception as e:
            log.warning(f"pdf page skip: {e}")
    return "\n".join(texts)


def parse_pagemaker(data: bytes) -> str:
    """Adobe PageMaker (.pmd/.p65/.pm6/.pm7) — OLE compound docs with UTF-16 text streams.
    Legacy files with embedded ChemDraw/MathType typically contain very little decodable
    body text — that's a limitation of the format, not the parser."""
    # Noise patterns to strip from extracted runs
    NOISE = [
        r"\{[0-9A-Fa-f\-]{8}(?:-[0-9A-Fa-f]{4}){3}-[0-9A-Fa-f]{12}\}",  # GUIDs
        r"CS ChemDraw Drawing.*?", r"MathType.*?Equation.*?", r"CorelDRAW.*?Graphic.*?",
        r"Equation Native", r"OlePres\d+", r"CompObj", r"CONTENTS",
        r"WinAllBasicCodePages", r"MT Extra", r"ChemDraw\.?[A-Za-z0-9\.]*",
        r"ChemDraw Interchange Format", r"content/[a-z0-9/\.]+", r"font/[a-z0-9/\.]+",
        r"META-INF/[a-z0-9/\.]+", r"mimetype[a-z\.\-\+/]+", r"application/[a-z\.\-\+/]+",
        r"MathType\S*", r"Ole10Native", r"Package",
    ]
    noise_re = re.compile("|".join(NOISE))

    def _clean_run(chunk: str) -> str:
        cleaned = noise_re.sub(" ", chunk).strip()
        # collapse runs of non-letter garbage
        cleaned = re.sub(r"[^\x20-\x7E\u00A0-\u024F\u2010-\u2E7F]+", " ", cleaned)
        cleaned = re.sub(r"\s{2,}", " ", cleaned).strip()
        return cleaned

    def _keep(cleaned: str) -> bool:
        if len(cleaned) < 8: return False
        letters = sum(1 for c in cleaned if c.isalpha())
        if letters < 5: return False
        # letter density > 40% AND at least one 4+ letter word
        if letters / max(len(cleaned), 1) < 0.4: return False
        if not re.search(r"[A-Za-z]{4,}", cleaned): return False
        return True

    text_parts: list[str] = []

    # 1) OLE walk (real .pmd files)
    if data[:8] == b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1":
        try:
            import olefile
            ole = olefile.OleFileIO(io.BytesIO(data))
            for stream_path in ole.listdir(streams=True):
                # skip embedded object binaries (never contain question text)
                if stream_path and stream_path[0] == "ObjectPool":
                    continue
                # skip metadata streams
                if any(s in ("\x01Ole", "\x01CompObj", "\x05SummaryInformation",
                              "\x05DocumentSummaryInformation") for s in stream_path):
                    continue
                try:
                    with ole.openstream(stream_path) as sf:
                        blob = sf.read()
                    # PageMaker stores runs mostly as UTF-16-LE; also try latin-1
                    for encoding in ("utf-16-le", "latin-1"):
                        try:
                            decoded = blob.decode(encoding, errors="ignore")
                        except Exception:
                            continue
                        # extract printable runs (10+ chars for stricter filter)
                        for chunk in re.findall(r"[\x20-\x7E\u00A0-\u024F]{10,}", decoded):
                            cleaned = _clean_run(chunk)
                            if _keep(cleaned):
                                text_parts.append(cleaned)
                except Exception as e:
                    log.debug(f"stream skip {stream_path}: {e}")
            ole.close()
        except Exception as e:
            log.warning(f"OLE parse failed: {e}")

    # 2) Fallback raw scan if OLE gave us very little
    if len(" ".join(text_parts)) < 500:
        raw = data.decode("latin-1", errors="ignore")
        for chunk in re.findall(r"[\x20-\x7E]{25,}", raw):
            cleaned = _clean_run(chunk)
            if _keep(cleaned):
                text_parts.append(cleaned)

    # dedupe, join
    seen = set()
    out: list[str] = []
    for t in text_parts:
        key = re.sub(r"\s+", " ", t.strip())
        if key and key not in seen and len(key) > 6:
            seen.add(key)
            out.append(key)
    joined = "\n".join(out)
    return joined[:80000]


async def download_google_drive(url: str) -> Tuple[bytes, str]:
    """Download a public Google Drive file. Accepts share links, /file/d/ID/, id= links."""
    m = re.search(r"/d/([a-zA-Z0-9_-]+)", url) or re.search(r"[?&]id=([a-zA-Z0-9_-]+)", url)
    if not m:
        raise ValueError("Could not extract Google Drive file id from URL")
    file_id = m.group(1)
    dl_url = f"https://drive.google.com/uc?export=download&id={file_id}"
    async with httpx.AsyncClient(follow_redirects=True, timeout=60) as c:
        r = await c.get(dl_url)
        if r.status_code != 200:
            raise ValueError(f"Drive download failed: {r.status_code}")
        # try to detect filename from content-disposition
        cd = r.headers.get("content-disposition", "")
        fname = re.search(r'filename="?([^"]+)"?', cd)
        return r.content, (fname.group(1) if fname else f"drive_{file_id}")
